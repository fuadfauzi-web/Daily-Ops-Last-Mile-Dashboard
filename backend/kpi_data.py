"""KPI Dashboard data uploads (staging, 2026-09-26).

Until the Metabase link works -- and for the KPI data that is still pasted by hand -- a manager / admin downloads the data
(from Metabase or the team's RCA sheets) and uploads it here. One current file per dataset: uploading again replaces it. The file
itself goes to object storage (storage.py); the database keeps who / when / how many rows. Rows are parsed on demand and cached in
memory until the upload changes.

Accepted: .csv, or an .xlsx (the dataset's `sheet` is picked, else the only sheet, else the sheet that has the required columns).
Column names are matched loosely (case, spaces and punctuation ignored), so "Courier Display Name" and "courier_display_name" both work.
"""
import asyncio
from functools import lru_cache
import csv
import io
import logging
import re
import sys
import uuid
from datetime import date, datetime, timezone

import xlsx_fast
from starlette.concurrency import run_in_threadpool

import db
import storage

log = logging.getLogger("kpi_data")

MAX_UPLOAD_BYTES = 40 * 1024 * 1024
MAX_ROWS = 400_000


# Metabase names a date column that is grouped by a time unit "Route Date: Day" / "Route Week: Month" -- the unit is not part of what the
# column means, so it is dropped before matching (2026-09-26: the Hybrid daily export was refused for "missing routedate").
_UNIT_SUFFIX = re.compile(r":\s*(minute|hour|day|week|month|quarter|year)\s*$", re.I)


@lru_cache(maxsize=8192)
def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", _UNIT_SUFFIX.sub("", text).lower())


# a column that goes by two names in Metabase ("Start Clock: Day" from the datetime, "Start Clock Date" from a date column) is one column
_ALIASES = {"startclockdate": "startclock"}


def norm(name) -> str:
    """Column names are normalised for every cell of every row on every request -- the same few dozen strings -- so it is cached."""
    n = _norm(str(name))
    return _ALIASES.get(n, n)


_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def to_iso_day(value) -> str:
    """A day as yyyy-mm-dd from what an export holds: an ISO date / timestamp (Metabase JSON, Excel cells), "September 25, 2026"
    (Metabase's formatted CSV), "25 Sep 2026" or 25/09/2026 (day first). "" when it is not a date."""
    s = str(value or "").strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.search(r"([A-Za-z]{3,9})\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})", s)  # September 25, 2026
    if m and m.group(1)[:3].lower() in _MONTHS:
        return f"{m.group(3)}-{_MONTHS[m.group(1)[:3].lower()]:02d}-{int(m.group(2)):02d}"
    m = re.search(r"(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3,9})\.?,?\s+(\d{4})", s)  # 25 Sep 2026
    if m and m.group(2)[:3].lower() in _MONTHS:
        return f"{m.group(3)}-{_MONTHS[m.group(2)[:3].lower()]:02d}-{int(m.group(1)):02d}"
    m = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", s)  # 25/09/2026 -- day first, like the team's sheets
    if m:
        d, mo = int(m.group(1)), int(m.group(2))
        if mo > 12 >= d:
            d, mo = mo, d
        if 1 <= mo <= 12 and 1 <= d <= 31:
            return f"{m.group(3)}-{mo:02d}-{d:02d}"
    return ""


# dataset -> what it is, which KPI page it feeds, which sheet of a workbook to read, and the columns it must have (normalised).
DATASETS: dict[str, dict] = {
    "hybrid_weekly": {
        "kpi": "hybrid", "label": "Hybrid weekly", "hint": "Metabase question 127194 (Hybrid Weekly Apps - All Regions) -- Download results as .csv",
        "link": "https://metabase.ninjavan.co/question/127194",
        "sheet": "Weekly Raw", "required": ["courierdisplayname", "routeweek"],
    },
    "hybrid_monthly": {
        "kpi": "hybrid", "label": "Hybrid monthly", "hint": "Metabase question 127195 (Hybrid Monthly Apps - All Regions) -- Download results as .csv",
        "link": "https://metabase.ninjavan.co/question/127195",
        "sheet": "Monthly Raw", "required": ["courierdisplayname", "routemonth"],
    },
    "hybrid_daily": {
        "kpi": "hybrid", "label": "Hybrid daily", "hint": "Metabase question 127196 (Hybrid Daily Apps - All Regions, current month) -- Download results as .csv",
        "link": "https://metabase.ninjavan.co/question/127196",
        "sheet": "Daily Raw", "required": ["courierdisplayname", "routedate"],
    },
    "hybrid_data": {
        "kpi": "hybrid", "label": "Hybrid driver list", "hint": "Metabase question 127193 (Hybrid Data Current Year - All Regions) -- Download results as .csv; gives each driver's start date (Service Duration)",
        "link": "https://metabase.ninjavan.co/question/127193",
        "sheet": "Hybrid Data", "required": ["displayname"],
    },
    "invalid_pod_raw": {
        "kpi": "invalid_pod", "label": "POD validation (raw)",
        "hint": "Metabase question 69573 (POP/POD Validation Tasks Raw Data): leave Hub Region empty for all regions, pick Date Type and Start / End date, then Download results as .csv -- or the Raw sheet of the POD Validation Analysis file",
        "link": "https://metabase.ninjavan.co/question/69573?transaction_type=DELIVERY&hub_region=&shipper_id=&date_type=&parent_id_coalesce=&start_date=&end_date=&driver_type=",
        "sheet": "Raw", "required": ["hubshortname", "validationresult"],
        "keep": ["hubshortname", "couriername", "trackingid", "transactionfailurereason", "validationresult", "invalidpodreason", "attempteddatetime", "validationdatetime", "validationusername"],
    },
    "pod_performance": {
        "kpi": "invalid_pod", "label": "LM POD performance (managers + admins)", "link": None,
        "hint": "the RAW DATA sheet of the LM POD Performance workbook (adds the audit result, final result / reason, zone and route type) -- only managers and admins see this view",
        "sheet": "RAW DATA", "required": ["hubshortname", "couriername", "result"],
        "keep": ["week", "hubshortname", "couriername", "transactionfailurereason", "validationdatetime", "validationresult", "invalidpodreason", "auditresult", "result", "finalreason", "hubzone", "routetype"],
    },
    "cod_rts_cod": {
        "kpi": "cod_rts", "label": "COD RTS (raw COD)", "link": "https://metabase.ninjavan.co/question/127198",
        "hint": "Metabase question 127198 (COD RTS Rate - All Regions, previous week): change the date filter for another week, then Download results as .csv -- or the RAW COD sheet of the RTS Analysis file",
        "sheet": "RAW COD", "required": ["desthubname", "rtsflag"],
    },
    "cod_rts_overall": {
        "kpi": "cod_rts", "label": "RTS overall (raw)", "link": "https://metabase.ninjavan.co/question/127197",
        "hint": "Metabase question 127197 (RTS Overall - All Regions, previous week): use the same week as the COD file, then Download results as .csv -- or the RAW Overal sheet; optional",
        "sheet": "RAW Overal", "required": ["trackingid", "rtsreason"],
    },
    "weekly_kpi": {
        "kpi": "weekly", "label": "Weekly KPI results", "hint": "the Station KPI W0W sheet of the Dashboard WoW file (or its CSV)",
        "sheet": "Station KPI W0W", "required": ["week", "station"],
    },
    "cisp_prior": {
        "kpi": "prior", "label": "Prior KPI (station by day)", "link": "https://metabase.ninjavan.co/question/127199",
        "hint": "Metabase question 127199 (CISP Prior - station by START CLOCK day, last 35 days) -- Download results as .csv. PRE-tagged TNs only, open PETs excluded; each TN is measured on its working start clock date, the result is by start clock date; provisional exclusions until OPEX confirms",
        "sheet": None, "required": ["desthubname", "startclock", "measured", "met"],
    },
    "cisp_completion": {
        "kpi": "completion", "label": "Completion D0 + D3 (station by day)", "link": "https://metabase.ninjavan.co/question/127200",
        "hint": "Metabase question 127200 (CISP Completion D0 + D3 - station by day, last 35 days) -- Download results as .csv; used by both Completion D0 and Completion D3. Provisional exclusions until OPEX confirms",
        "sheet": None, "required": ["desthubname", "lastmilestartclockdate", "d0measured", "d0met", "d3measured", "d3met"],
    },
    "cisp_terminal": {
        "kpi": "terminal", "label": "Terminal T7 (station by day)", "link": "https://metabase.ninjavan.co/question/127201",
        "hint": "Metabase question 127201 (CISP Terminal T7 - station by day, TNs past their N7 cut-off in the last 35 days) -- Download results as .csv. Provisional exclusions until OPEX confirms",
        "sheet": None, "required": ["desthubname", "lastmilestartclockdate", "t7measured", "t7met"],
    },
    "cisp_fifo": {
        "kpi": "fifo", "label": "FIFO D0 (per station, one period)", "link": "https://metabase.ninjavan.co/question/118041-my-lm-cisp-fifo-d0?FIFO_target=92&measured_date=past1weeks",
        "hint": "Metabase question 118041 ((MY) LM CISP FIFO D0): pick the Measured Date (e.g. past 1 weeks), run it, then Download results as .csv -- one period per file",
        "sheet": None, "required": ["desthubname", "totalorders", "totaln0met"],
    },
    "region_list": {
        "kpi": "region", "label": "Station list (Region List sheet)", "link": "https://docs.google.com/spreadsheets/d/1KmHiK5q5mMoKX8N2TzlmRByjIc5nX2fwSsuCxHm4l8g/edit?gid=1339991125#gid=1339991125",
        "link_label": "Open the Region List sheet",
        "hint": "the Region List sheet (Region tab) downloaded as .csv or .xlsx -- Active / Virtual stations in Klang Valley, Northern, Southern, East Coast and East Malaysia become the station list; Closed, SAMEDAY and NO HUB rows are left out",
        "sheet": None, "required": ["grouplh", "stationname", "zone", "region"],
    },
    "opex_result": {
        "kpi": "opex", "label": "OPEX dashboard result", "link": "https://last-mile-dashboard.ninjavan.apps.substrait.build/",
        "hint": "the OPEX Last Mile Performance dashboard: pick the region / area and the dates, press Download CSV, then upload that file (any other table is shown as it is)",
        "sheet": None, "required": [],
    },
}
KPI_DATASETS = {k: [d for d, v in DATASETS.items() if v["kpi"] == k] for k in {v["kpi"] for v in DATASETS.values()}}


class UploadError(ValueError):
    pass


# ------------------------------------------------------------------------------------------------ parsing

def _clean_header(value, idx: int) -> str:
    s = " ".join(str(value).split()) if value is not None else ""
    return s or f"_col{idx}"


def _cell(v):
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    if isinstance(v, str) and len(v) < 80:
        return sys.intern(v)  # hub / reason / driver names repeat on thousands of rows -- one copy in memory
    return v


def _keep_fn(keep):
    if not keep:
        return None
    wanted = set(keep)
    return lambda header_text: norm(header_text) in wanted


def _parse_csv(data: bytes, keep=None) -> tuple[list[str], list[dict]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    reader = csv.reader(io.StringIO(text))
    header = None
    positions: list[tuple[int, str]] = []
    keep_fn = _keep_fn(keep)
    rows: list[dict] = []
    for raw in reader:
        if header is None:
            if sum(1 for c in raw if str(c).strip()) >= 2 or len(raw) == 1 and str(raw[0]).strip():
                full = [_clean_header(c, i) for i, c in enumerate(raw)]
                positions = [(i, h) for i, h in enumerate(full) if keep_fn is None or keep_fn(h)]
                header = [h for _i, h in positions]
            continue
        if not any(str(c).strip() for c in raw):
            continue
        rows.append({h: (sys.intern(raw[i]) if i < len(raw) and len(raw[i]) < 80 else (raw[i] if i < len(raw) else "")) for i, h in positions})
        if len(rows) > MAX_ROWS:
            raise UploadError(f"More than {MAX_ROWS:,} rows -- split the file")
    if header is None:
        raise UploadError("The file is empty")
    return header, rows


def _pick_sheet(names: list[str], wanted: str | None, has_required) -> str | None:
    if wanted:
        for n in names:
            if n.strip().lower() == wanted.lower():
                return n
        for n in names:
            if n.strip().lower().startswith(wanted.lower()[:8]):
                return n
    if len(names) == 1:
        return names[0]
    return None


def _parse_xlsx_fast(data: bytes, sheet: str | None, required: list[str], keep_fn) -> tuple[list[str], list[dict]]:
    """Same sheet choice as the openpyxl path below, read with xlsx_fast (about 4x quicker, only the wanted columns)."""
    names = xlsx_fast.sheet_names(data)
    chosen = _pick_sheet(names, sheet, None)
    candidates = [chosen] if chosen else names

    def read(name: str):
        try:
            return xlsx_fast.read_sheet(data, name, keep_fn, _clean_header, _cell, MAX_ROWS)
        except ValueError as exc:
            if "too many rows" in str(exc):
                raise UploadError(f"More than {MAX_ROWS:,} rows -- split the file") from exc
            raise

    for name in candidates:
        got = read(name)
        if got and all(r in {norm(h) for h in got[0]} for r in required):
            return got
    if chosen is None and required:
        raise UploadError(f"Couldn't find a sheet with the columns this needs. Sheets in the file: {', '.join(names)}")
    got = read(chosen) if chosen else None
    if got:
        return got
    raise UploadError("The sheet is empty")


def _parse_xlsx(data: bytes, sheet: str | None, required: list[str], keep=None) -> tuple[list[str], list[dict]]:
    try:
        return _parse_xlsx_fast(data, sheet, required, _keep_fn(keep))
    except UploadError:
        raise
    except Exception:  # noqa: BLE001 - an unusual workbook: fall back to openpyxl (slower, reads everything)
        log.warning("Fast xlsx reader failed -- falling back to openpyxl", exc_info=True)
    import openpyxl  # heavy import, only when the fast reader could not handle the file

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        names = wb.sheetnames

        def read(name: str) -> tuple[list[str], list[dict]] | None:
            it = wb[name].iter_rows(values_only=True)
            header = None
            rows: list[dict] = []
            for i, raw in enumerate(it):
                if header is None:
                    if sum(1 for c in raw if c is not None and str(c).strip()) >= 2:
                        header = [_clean_header(c, j) for j, c in enumerate(raw)]
                    elif i > 30:
                        return None
                    continue
                if not any(c is not None and str(c).strip() for c in raw):
                    continue
                rows.append({h: _cell(raw[j]) if j < len(raw) else None for j, h in enumerate(header)})
                if len(rows) > MAX_ROWS:
                    raise UploadError(f"More than {MAX_ROWS:,} rows -- split the file")
            return (header, rows) if header else None

        chosen = _pick_sheet(names, sheet, None)
        candidates = [chosen] if chosen else names
        for name in candidates:
            got = read(name)
            if got and all(r in {norm(h) for h in got[0]} for r in required):
                return got
        if chosen is None and required:
            raise UploadError(f"Couldn't find a sheet with the columns this needs. Sheets in the file: {', '.join(names)}")
        got = read(chosen) if chosen else None
        if got:
            return got
        raise UploadError("The sheet is empty")
    finally:
        wb.close()


def parse_table(filename: str, data: bytes, sheet: str | None, required: list[str], keep=None) -> tuple[list[str], list[dict]]:
    """keep = normalised column names to retain (the required ones are always kept) -- a big file only costs memory for what is used."""
    name = (filename or "").lower()
    keep = list(dict.fromkeys([*keep, *required])) if keep else None
    if name.endswith(".csv"):
        header, rows = _parse_csv(data, keep)
    elif name.endswith((".xlsx", ".xlsm")):
        header, rows = _parse_xlsx(data, sheet, required, keep)
    else:
        raise UploadError("Upload a .csv or .xlsx file")
    have = {norm(h) for h in header}
    missing = [r for r in required if r not in have]
    if missing:
        raise UploadError(f"This doesn't look like the right file -- these columns are missing: {', '.join(missing)}. Found: {', '.join(header[:12])}")
    if not rows:
        raise UploadError("The file has a header but no data rows")
    return header, rows


# ------------------------------------------------------------------------------------------------ storage + cache

_cache: dict[str, dict] = {}  # dataset -> {"stamp": "yyyy-mm-dd hh:mm:ss", "rows": [...], "header": [...]}


def _stamp(value) -> str:
    """Upload time as a comparable string (DB values are naive UTC, ours are aware UTC -- the digits are the same)."""
    return str(value)[:19].replace("T", " ")
_locks: dict[str, asyncio.Lock] = {}


def _lock(dataset: str) -> asyncio.Lock:
    return _locks.setdefault(dataset, asyncio.Lock())


async def list_uploads() -> dict[str, dict]:
    rows = await db.fetch_all("SELECT dataset, filename, row_count, uploaded_by, uploaded_at FROM kpi_uploads")
    return {
        r[0]: {"filename": r[1], "row_count": r[2], "uploaded_by": r[3], "uploaded_at": r[4].isoformat() if hasattr(r[4], "isoformat") else str(r[4])}
        for r in rows
    }


async def save_upload(dataset: str, filename: str, data: bytes, user_email: str) -> dict:
    spec = DATASETS.get(dataset)
    if spec is None:
        raise UploadError("Unknown dataset")
    if not data:
        raise UploadError("The file is empty")
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError(f"The file is too large (max {MAX_UPLOAD_BYTES // (1024 * 1024)} MB)")
    header, rows = await run_in_threadpool(parse_table, filename, data, spec["sheet"], spec["required"], spec.get("keep"))
    if spec.get("check"):  # a dataset may refuse a file that has the columns but is not the right one (the Region List: too few stations)
        await run_in_threadpool(spec["check"], rows)
    ext = ".xlsx" if filename.lower().endswith((".xlsx", ".xlsm")) else ".csv"
    key = storage.safe_key("kpi", f"{dataset}-{uuid.uuid4().hex}{ext}")
    content_type = "application/octet-stream" if ext == ".xlsx" else "text/csv"
    await run_in_threadpool(storage.put_bytes, key, data, content_type=content_type)
    old = await db.fetch_one("SELECT storage_key FROM kpi_uploads WHERE dataset = %s", (dataset,))
    now = datetime.now(timezone.utc).replace(microsecond=0)
    await db.execute("DELETE FROM kpi_uploads WHERE dataset = %s", (dataset,))
    await db.execute(
        "INSERT INTO kpi_uploads (dataset, storage_key, filename, row_count, uploaded_by, uploaded_at) VALUES (%s, %s, %s, %s, %s, %s)",
        (dataset, key, filename[:255], len(rows), user_email, now),
    )
    if old and old[0]:
        try:
            await run_in_threadpool(storage.delete, old[0])
        except Exception:  # noqa: BLE001 - a leftover blob is harmless
            log.warning("Could not delete the previous %s upload", dataset)
    builder = _builders.get(dataset)
    if builder is not None:
        # big datasets keep only a compact, pre-aggregated structure in memory (built once per upload), not thousands of row dicts
        compact = await run_in_threadpool(builder, rows)
        async with _lock(dataset):
            _compact[dataset] = {"stamp": _stamp(now), "data": compact}
    else:
        async with _lock(dataset):
            _cache[dataset] = {"stamp": _stamp(now), "rows": rows, "header": header}
    return {"row_count": len(rows), "columns": header}


async def delete_upload(dataset: str) -> None:
    old = await db.fetch_one("SELECT storage_key FROM kpi_uploads WHERE dataset = %s", (dataset,))
    await db.execute("DELETE FROM kpi_uploads WHERE dataset = %s", (dataset,))
    _cache.pop(dataset, None)
    _compact.pop(dataset, None)
    if old and old[0]:
        try:
            await run_in_threadpool(storage.delete, old[0])
        except Exception:  # noqa: BLE001
            log.warning("Could not delete the %s upload file", dataset)


_compact: dict[str, dict] = {}  # dataset -> {"stamp", "data"}: the compact structure a dataset's builder made from its rows
_builders: dict[str, object] = {}


def drop_hub_caches(keep: tuple = ()) -> None:
    """The compact structures carry station names / zones / regions taken from the station list when they were built: when the station list changes they are
    dropped and rebuilt from the stored file on the next request."""
    for name in [n for n in _compact if n not in keep]:
        _compact.pop(name, None)


def register_compact(dataset: str, builder) -> None:
    """A dataset with a compact builder never keeps its row dicts: rows -> builder(rows) once per upload, and only that is cached."""
    _builders[dataset] = builder


async def load_compact(dataset: str) -> tuple[dict, object] | None:
    """(meta, compact structure) of the current upload, parsed and built once per upload; None if nothing was uploaded."""
    meta_row = await db.fetch_one("SELECT storage_key, filename, row_count, uploaded_by, uploaded_at FROM kpi_uploads WHERE dataset = %s", (dataset,))
    if meta_row is None:
        _compact.pop(dataset, None)
        return None
    key, filename, row_count, uploaded_by, uploaded_at = meta_row
    stamp = _stamp(uploaded_at)
    meta = {"filename": filename, "row_count": row_count, "uploaded_by": uploaded_by, "uploaded_at": uploaded_at.isoformat() if hasattr(uploaded_at, "isoformat") else str(uploaded_at)}
    async with _lock(dataset):
        hit = _compact.get(dataset)
        if hit and hit["stamp"] == stamp:
            return meta, hit["data"]
        spec = DATASETS[dataset]
        data = await run_in_threadpool(storage.get_bytes, key)
        _header, rows = await run_in_threadpool(parse_table, filename, data, spec["sheet"], spec["required"], spec.get("keep"))
        compact = await run_in_threadpool(_builders[dataset], rows)
        _compact[dataset] = {"stamp": stamp, "data": compact}
        return meta, compact


async def load_rows(dataset: str) -> tuple[dict, list[dict]] | None:
    """(meta, rows) of the current upload for a dataset, or None if nothing was uploaded."""
    meta_row = await db.fetch_one("SELECT storage_key, filename, row_count, uploaded_by, uploaded_at FROM kpi_uploads WHERE dataset = %s", (dataset,))
    if meta_row is None:
        _cache.pop(dataset, None)
        return None
    key, filename, row_count, uploaded_by, uploaded_at = meta_row
    stamp = _stamp(uploaded_at)
    meta = {"filename": filename, "row_count": row_count, "uploaded_by": uploaded_by, "uploaded_at": uploaded_at.isoformat() if hasattr(uploaded_at, "isoformat") else str(uploaded_at)}
    async with _lock(dataset):
        hit = _cache.get(dataset)
        if hit and hit["stamp"] == stamp:
            return meta, hit["rows"]
        spec = DATASETS[dataset]
        data = await run_in_threadpool(storage.get_bytes, key)
        header, rows = await run_in_threadpool(parse_table, filename, data, spec["sheet"], spec["required"], spec.get("keep"))
        _cache[dataset] = {"stamp": stamp, "rows": rows, "header": header}
        return meta, rows
